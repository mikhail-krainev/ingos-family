"""Сборка Финансовая_модель_ИНГОССЕМЬЯ.xlsx: вся финансовая модель «ИНГОССЕМЬЯ» на формулах Excel.

Исходные данные берутся из model.py (единый источник параметров). Все расчеты, кроме
анализа чувствительности и Монте-Карло, выполняются формулами в книге и пересчитываются
при изменении допущений. Чувствительность и Монте-Карло требуют тысяч прогонов модели,
поэтому их результаты переносятся из model.py значениями.

Запуск: вызывается из model.py (python model.py).
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

import model as M

FONT_NAME = "Times New Roman"
FONT_SIZE = 11
INPUT_COLOR = "0000FF"
LINK_COLOR = "008000"

F_RUB = '#,##0;(#,##0);"-"'
F_NUM = '#,##0;(#,##0);"-"'
F_NUM1 = '#,##0.0;(#,##0.0);"-"'
F_DEC = '0.00;(0.00);"-"'
F_DEC3 = '0.000'
F_PCT = '0.0%;(0.0%);"-"'
F_INT = '0'

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SCEN_COLS = {"Пессимистичный": "C", "Базовый": "D", "Оптимистичный": "E", "Версия 1": "F"}
SCEN_OVERRIDES = {**M.SCENARIOS, "Версия 1": {k: v for k, v in M.V1_MANUAL.items() if k != "fixed_load"}}
A_SHEET = "Допущения"
A_Q = "'Допущения'"

# Подписи и единицы для параметров model.py: ключ -> (название, единица, формат)
LABELS = {
    "vat": ("НДС", "доля", F_PCT),
    "profit_tax": ("Налог на прибыль", "доля", F_PCT),
    "loss_carry_limit": ("Лимит переноса убытков от налоговой базы года", "доля", F_PCT),
    "payroll_mult": ("Коэффициент страховых взносов к ФОТ", "коэф.", F_DEC3),
    "discount_rate_y": ("Ставка дисконтирования, годовая", "доля", F_PCT),
    "salary_index_y": ("Индексация зарплат в год", "доля", F_PCT),
    "price_index_y": ("Индексация цен подписки в год, со 2-го года", "доля", F_PCT),
    "acquiring": ("Эквайринг от суммы оплаты", "доля", F_PCT),
    "overhead_share": ("Накладные расходы от ФОТ", "доля", F_PCT),
    "price_basic": ("Цена тарифа «Базовый» с НДС", "руб./мес.", F_RUB),
    "price_ext": ("Цена тарифа «Расширенный» с НДС", "руб./мес.", F_RUB),
    "share_ext": ("Доля тарифа «Расширенный» в B2C", "доля", F_PCT),
    "partner_spend": ("Покупки услуг партнеров через координатора", "руб./семью в мес.", F_RUB),
    "partner_commission": ("Агентская комиссия с услуг партнеров", "доля", F_PCT),
    "b2b_price": ("Цена опции к корпоративному ДМС с НДС", "руб./мес.", F_RUB),
    "int_cac": ("CAC во внутренних каналах", "руб.", F_RUB),
    "int_cap": ("Емкость внутренних каналов", "семей/мес.", F_NUM),
    "ext_cac": ("CAC во внешних каналах", "руб.", F_RUB),
    "mkt_start": ("Маркетинговый бюджет в 4-м месяце", "руб.", F_RUB),
    "mkt_growth_m": ("Рост маркетингового бюджета в месяц", "доля", F_PCT),
    "mkt_cap": ("Потолок маркетингового бюджета", "руб./мес.", F_RUB),
    "vet_start_month": ("Месяц запуска партнерского канала с фондом «Защитники Отечества»", "месяц", F_INT),
    "vet_new": ("Новых семей в месяц через партнерский канал фонда", "семей/мес.", F_NUM),
    "vet_cac": ("Затраты на привлечение в партнерском канале", "руб.", F_RUB),
    "b2b_start_month": ("Месяц первого корпоративного клиента", "месяц", F_INT),
    "b2b_clients_per_month": ("Новых корпоративных клиентов в месяц", "клиентов", F_DEC),
    "b2b_employees": ("Сотрудников у корпоративного клиента", "чел.", F_NUM),
    "b2b_activation": ("Доля сотрудников, подключивших родителя", "доля", F_PCT),
    "b2b_ramp_months": ("Срок разгона активации у клиента", "мес.", F_INT),
    "churn_first": ("Отток после первого месяца", "доля", F_PCT),
    "churn_b2c": ("Отток B2C в месяц со второго месяца", "доля", F_PCT),
    "churn_b2b": ("Отток B2B2C в месяц", "доля", F_PCT),
    "time_mult": ("Множитель к бюджету минут координатора", "коэф.", F_DEC),
    "coord_salary": ("Зарплата координатора", "руб./мес.", F_RUB),
    "supervisor_ratio": ("Координаторов на одного супервизора", "чел.", F_NUM),
    "supervisor_salary": ("Зарплата супервизора", "руб./мес.", F_RUB),
    "coord_turnover_y": ("Текучесть координаторов в год", "доля", F_PCT),
    "onboarding_cost": ("Стоимость обучения нового координатора", "руб.", F_RUB),
    "seat_cost": ("CRM и телефония на рабочее место", "руб./мес.", F_RUB),
    "telemed_per_family": ("Телемедицина на семью", "руб./мес.", F_RUB),
    "telecom_per_family": ("Связь на семью (робот, звонки, СМС)", "руб./мес.", F_RUB),
    "hosting_base": ("Размещение в контуре ИСПДн группы", "руб./мес.", F_RUB),
    "hosting_per_1000": ("Доплата за размещение на 1000 семей", "руб./мес.", F_RUB),
    "legal_entity_m": ("Сопровождение сервисной компании", "руб./мес.", F_RUB),
    "cross_rate_m": ("Доля семей, покупающих ДМС или полис ухода в месяц", "доля", '0.0%'),
    "cross_margin": ("Маржинальный доход группы с кросс-продажи", "руб.", F_RUB),
    "mvp_paying_share": ("Доля платящих семей в MVP", "доля", F_PCT),
}

# Параметры, которых нет в словарях model.py в явном виде
EXTRA_GENERAL = [
    ("work_hours", "Рабочих часов координатора в месяц", "ч", F_NUM, 165, "допущение", "Производственный календарь"),
    ("productive_share", "Доля времени на работу с семьями", "доля", F_PCT, 0.75, "допущение",
     "Остальное: обучение, планерки, документы"),
    ("coverage_mult", "Множитель покрытия смен", "коэф.", F_DEC, M.COVERAGE_MULT[0], M.COVERAGE_MULT[1], M.COVERAGE_MULT[2]),
    ("min_coords", "Минимальное число координаторов", "чел.", F_NUM, 3, "допущение", "Команда MVP"),
    ("mvp_months", "Длительность MVP", "мес.", F_INT, M.MVP_MONTHS, "условие кейса", "Первый MVP за 3 месяца"),
    ("mvp_limit", "Лимит бюджета MVP", "руб.", F_RUB, M.MVP_LIMIT, "условие кейса", "Не более 5 млн руб."),
    ("contingency", "Резерв на непредвиденные расходы MVP", "доля", F_PCT, M.MVP_CONTINGENCY[0], M.MVP_CONTINGENCY[1],
     M.MVP_CONTINGENCY[2]),
]
assert abs(165 * 60 * 0.75 - M.PRODUCTIVE_MIN[0]) < 1e-9

FIXED_LOAD = [
    ("fixed_load", "Фиксированная нагрузка, семей «Базовый»", "семей", F_NUM),
    ("fixed_ext_weight", "Вес семьи «Расширенный» к «Базовому»", "коэф.", F_DEC),
]

# Бюджет минут координатора: тариф, операция, раз в месяц, минут на операцию
TIME_ROWS = [
    ("Базовый", "Плановый звонок родителю", 1, 12),
    ("Базовый", "Перезвон по флагу голосового робота", 0.8, 10),
    ("Базовый", "Организация услуги (запись, анализы, доставка)", 1.2, 15),
    ("Базовый", "Проверка отчета ребенку (черновик готовит ИИ)", 1.2, 3),
    ("Базовый", "Ответы ребенку в чате", 1, 5),
    ("Расширенный", "Плановые звонки родителю", 4.3, 12),
    ("Расширенный", "Перезвон по флагу голосового робота", 1.2, 10),
    ("Расширенный", "Организация услуг", 2.5, 15),
    ("Расширенный", "Проверка отчетов ребенку", 2.5, 3),
    ("Расширенный", "Ответы ребенку в чате", 1, 10),
]
VETERAN_OP = "Помощь с льготами ветеранов (доля семей с льготой)"
for plan in ("Базовый", "Расширенный"):
    total = sum(n * t for p, _, n, t in TIME_ROWS if p == plan) + M.veteran_share() * M.VETERAN_PAPERWORK_MIN[0]
    assert abs(total - M.minutes(plan)) < 1e-9

TEAM_NOTES = {"Разработчик 2": "Второй разработчик под интеграции"}
CAPEX_NOTES = {4: "Интеграция в IngoMobile", 7: "Интеграция с клиниками через МедРокет",
               10: "ИИ-черновики отчетов и доработка робота"}


# ---------------------------------------------------------------------------
# Оформление
# ---------------------------------------------------------------------------

def font(bold=False, color="000000", size=FONT_SIZE, italic=False):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic)


def put(ws, ref, value, fmt=None, bold=False, color="000000", align=None, border=True, wrap=False, italic=False):
    c = ws[ref]
    c.value = value
    c.font = font(bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def title(ws, text, subtitle=None):
    ws["A1"].value = text
    ws["A1"].font = font(bold=True, size=14)
    if subtitle:
        ws["A2"].value = subtitle
        ws["A2"].font = font(italic=True)


def header(ws, row, names, col=1, height=None):
    for i, name in enumerate(names):
        put(ws, f"{get_column_letter(col + i)}{row}", name, bold=True, align="center", wrap=True)
    if height:
        ws.row_dimensions[row].height = height


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def page_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Лист &P из &N"
    ws.oddFooter.center.font = FONT_NAME
    ws.sheet_view.showGridLines = False


def q(sheet):
    return f"'{sheet}'"


# ---------------------------------------------------------------------------
# Лист «Допущения»
# ---------------------------------------------------------------------------

class Params:
    """Адреса параметров на листе «Допущения»."""

    def __init__(self):
        self.rows = {}

    def ref(self, key, scen):
        return f"{A_Q}!${SCEN_COLS[scen]}${self.rows[key]}"

    def fixed(self, key):
        return f"{A_Q}!$D${self.rows[key]}"


def build_assumptions(wb, P):
    ws = wb.create_sheet(A_SHEET)
    title(ws, "Исходные данные и допущения модели",
          "Синий шрифт: исходное значение, его можно менять. Черный шрифт: формула. "
          "Столбцы сценариев без собственного значения ссылаются на базовый сценарий.")
    widths(ws, [52, 16, 15, 15, 15, 15, 14, 70])
    r = 4
    header(ws, r, ["Параметр", "Единица", "Пессими-стичный", "Базовый", "Оптими-стичный", "Версия 1 без триажа",
                   "Метка", "Источник или обоснование"], height=32)

    def section(text):
        nonlocal r
        r += 1
        ws.merge_cells(f"A{r}:H{r}")
        put(ws, f"A{r}", text, bold=True)
        for col in "BCDEFGH":
            ws[f"{col}{r}"].border = BOX

    def param_row(key, label, unit, fmt, base_value, tag, note, overrides):
        nonlocal r
        r += 1
        P.rows[key] = r
        put(ws, f"A{r}", label, wrap=True)
        put(ws, f"B{r}", unit, align="center")
        put(ws, f"D{r}", base_value, fmt, color=INPUT_COLOR)
        for scen, col in SCEN_COLS.items():
            if scen == "Базовый":
                continue
            ov = overrides.get(scen, {})
            if key in ov:
                put(ws, f"{col}{r}", ov[key], fmt, color=INPUT_COLOR)
            else:
                put(ws, f"{col}{r}", f"=$D{r}", fmt)
        put(ws, f"G{r}", tag, align="center")
        put(ws, f"H{r}", note, wrap=True)

    section("1. Общие параметры")
    for key, (val, tag, note) in M.GENERAL.items():
        label, unit, fmt = LABELS[key]
        param_row(key, label, unit, fmt, val, tag, note, {})
    for key, label, unit, fmt, val, tag, note in EXTRA_GENERAL:
        param_row(key, label, unit, fmt, val, tag, note, {})

    section("2. Параметры сценариев")
    for key, (val, tag, note) in M.BASE.items():
        if isinstance(val, dict):
            continue
        label, unit, fmt = LABELS[key]
        param_row(key, label, unit, fmt, val, tag, note, SCEN_OVERRIDES)

    section("3. Нагрузка координатора в версии 1 (без голосового робота)")
    load, weight = M.V1_MANUAL["fixed_load"]
    for (key, label, unit, fmt), v1 in zip(FIXED_LOAD, (load, weight)):
        r += 1
        P.rows[key] = r
        put(ws, f"A{r}", label, wrap=True)
        put(ws, f"B{r}", unit, align="center")
        for scen, col in SCEN_COLS.items():
            put(ws, f"{col}{r}", v1 if scen == "Версия 1" else None, fmt, color=INPUT_COLOR)
        put(ws, f"G{r}", "допущение", align="center")
        put(ws, f"H{r}", "Координатор звонит каждой семье раз в неделю; в сценариях 1-3 нагрузка считается от минут", wrap=True)

    section("4. Расчетные параметры")
    t_sheet = q("Время координатора")
    derived = [
        ("min_basic", "Минут координатора на семью «Базовый» с множителем", "мин./мес.", F_NUM1,
         lambda c: f"={t_sheet}!$E${TIME_TOTALS['Базовый']}*{c}{P.rows['time_mult']}", "Лист «Время координатора»"),
        ("min_ext", "Минут координатора на семью «Расширенный» с множителем", "мин./мес.", F_NUM1,
         lambda c: f"={t_sheet}!$E${TIME_TOTALS['Расширенный']}*{c}{P.rows['time_mult']}", "Лист «Время координатора»"),
        ("capacity", "Продуктивных минут координатора с учетом покрытия смен", "мин./мес.", F_NUM,
         lambda c: f"={c}{P.rows['work_hours']}*60*{c}{P.rows['productive_share']}/{c}{P.rows['coverage_mult']}",
         "Часы x 60 x доля / множитель покрытия"),
        ("load_b", "Семей «Базовый» на координатора", "семей", F_NUM1,
         lambda c: f'=IF({c}{P.rows["fixed_load"]}="",{c}{{capacity}}/{c}{{min_basic}},'
                   f'{c}{P.rows["fixed_load"]}/{c}{P.rows["coverage_mult"]})',
         "Емкость / минуты на семью"),
        ("load_e", "Семей «Расширенный» на координатора", "семей", F_NUM1,
         lambda c: f'=IF({c}{P.rows["fixed_load"]}="",{c}{{capacity}}/{c}{{min_ext}},'
                   f'{c}{P.rows["fixed_load"]}/{c}{P.rows["fixed_ext_weight"]}/{c}{P.rows["coverage_mult"]})',
         "Емкость / минуты на семью"),
        ("arpu_list", "Средний чек B2C с НДС по прейскуранту, 1-й год", "руб./мес.", F_RUB,
         lambda c: f"={c}{P.rows['price_basic']}*(1-{c}{P.rows['share_ext']})+{c}{P.rows['price_ext']}*{c}{P.rows['share_ext']}",
         "Взвешенная цена тарифов"),
        ("arpu_gross", "Средний чек B2C с НДС с учетом льгот ветеранам, 1-й год", "руб./мес.", F_RUB,
         lambda c: f"={c}{{arpu_list}}*(1-{VET['discount']})", "Лист «Льготы ветеранам»"),
        ("vet_k", "Недополученная выручка на 1 руб. оплаченной подписки", "коэф.", '0.0000',
         lambda c: f"={VET['discount']}/(1-{VET['discount']})", "Скидка / (1 - скидка)"),
        ("cohort_b2b", "Семей B2B2C от клиентов одного месяца", "семей", F_NUM1,
         lambda c: f"={c}{P.rows['b2b_clients_per_month']}*{c}{P.rows['b2b_employees']}*{c}{P.rows['b2b_activation']}",
         "Клиенты x сотрудники x активация"),
        ("disc_m", "Ставка дисконтирования, месячная", "доля", '0.000%',
         lambda c: f"=(1+{c}{P.rows['discount_rate_y']})^(1/12)-1", "Эквивалент годовой ставки"),
    ]
    for key, *_ in derived:
        r += 1
        P.rows[key] = r
    r_start = r - len(derived)
    for i, (key, label, unit, fmt, fn, note) in enumerate(derived, start=1):
        row = r_start + i
        put(ws, f"A{row}", label, wrap=True)
        put(ws, f"B{row}", unit, align="center")
        for col in SCEN_COLS.values():
            formula = fn(col).replace("{capacity}", str(P.rows["capacity"])).replace(
                "{min_basic}", str(P.rows["min_basic"])).replace("{min_ext}", str(P.rows["min_ext"])).replace(
                "{arpu_list}", str(P.rows["arpu_list"]))
            put(ws, f"{col}{row}", formula, fmt)
        put(ws, f"G{row}", "расчет", align="center")
        put(ws, f"H{row}", note, wrap=True)

    # Таблицы, общие для всех сценариев
    r += 2
    put(ws, f"A{r}", "5. Семьи в MVP", bold=True, border=False)
    r += 1
    header(ws, r, ["Месяц MVP", "Семей"])
    P.mvp_fam = {}
    for m, fam in M.BASE["mvp_families"][0].items():
        r += 1
        P.mvp_fam[m] = r
        put(ws, f"A{r}", m, F_INT, align="center")
        put(ws, f"B{r}", fam, F_NUM, color=INPUT_COLOR)
    P.mvp_fam[0] = None

    r += 2
    put(ws, f"A{r}", "6. Постоянная команда после MVP", bold=True, border=False)
    r += 1
    header(ws, r, ["Роль", "Ставка, FTE", "Зарплата, руб./мес.", "Месяц выхода", "Комментарий"])
    team = M.BASE["team"][0]
    P.team_first = r + 1
    for role, (fte, sal, start) in team.items():
        r += 1
        put(ws, f"A{r}", role)
        put(ws, f"B{r}", fte, F_DEC, color=INPUT_COLOR)
        put(ws, f"C{r}", sal, F_RUB, color=INPUT_COLOR)
        put(ws, f"D{r}", start, F_INT, color=INPUT_COLOR, align="center")
        put(ws, f"E{r}", TEAM_NOTES.get(role, ""))
    P.team_last = r

    r += 2
    put(ws, f"A{r}", "7. Инвестиции после MVP", bold=True, border=False)
    r += 1
    header(ws, r, ["Назначение", "Месяц", "Сумма, руб."])
    capex = M.BASE["capex"][0]
    P.capex_first = r + 1
    note = ""
    for m, amount in capex.items():
        r += 1
        note = CAPEX_NOTES.get(m, note)
        put(ws, f"A{r}", note)
        put(ws, f"B{r}", m, F_INT, align="center", color=INPUT_COLOR)
        put(ws, f"C{r}", amount, F_RUB, color=INPUT_COLOR)
    P.capex_last = r
    r += 1
    put(ws, f"A{r}", "Итого", bold=True)
    put(ws, f"B{r}", None)
    put(ws, f"C{r}", f"=SUM(C{P.capex_first}:C{P.capex_last})", F_RUB, bold=True)

    ws.freeze_panes = "C5"
    page_setup(ws)


# ---------------------------------------------------------------------------
# Лист «Льготы ветеранам»
# ---------------------------------------------------------------------------

VET = {}
VET_SHEET = "Льготы ветеранам"
VET_CONDITIONS = [
    ("Цена", "Скидки по категориям из таблицы выше; для родителей погибших участников СВО тариф «Базовый» бесплатно"),
    ("Навигатор льгот", "Координатор помогает получить положенное по закону: внеочередную медпомощь, льготные лекарства, "
                        "санаторно-курортное лечение, меры региональной поддержки"),
    ("Приоритет", "Отдельная линия координаторов, прошедших курс по льготам ветеранов и работе с семьями участников СВО"),
    ("Психологическая поддержка", "Консультация психолога через телемедицину Ингосстраха для семей участников СВО"),
    ("Партнерство", "Обмен обращениями с фондом «Защитники Отечества»: социальный координатор фонда передает задачи "
                    "по медпомощи родителю, ИНГОССЕМЬЯ сообщает фонду о социальных вопросах семьи"),
    ("Работа для ветеранов", "Ветераны СВО и члены семей участников СВО в приоритете при найме координаторов"),
    ("Проверка статуса", "Удостоверение проверяется при подключении; в CRM хранится только категория льготы и дата проверки"),
]


def build_veterans(wb):
    ws = wb.create_sheet(VET_SHEET)
    title(ws, "Особые условия для ветеранов СВО и ветеранов",
          "Льгота действует на подписку B2C. Скидка уменьшает выручку во всех сценариях; ее сумма показана на листе «Сводка».")
    widths(ws, [52, 14, 12, 14, 70])
    r = 4
    header(ws, r, ["Категория", "Доля семей B2C", "Скидка с подписки", "Вклад в среднюю скидку", "Условие и обоснование"],
           height=48)
    first = r + 1
    for cat, share, disc, tag, note in M.VETERAN_BENEFITS:
        r += 1
        put(ws, f"A{r}", cat, wrap=True)
        put(ws, f"B{r}", share, F_PCT, color=INPUT_COLOR)
        put(ws, f"C{r}", disc, '0%', color=INPUT_COLOR)
        put(ws, f"D{r}", f"=B{r}*C{r}", '0.00%')
        put(ws, f"E{r}", f"{note} [{tag}]", wrap=True)
    last = r
    r += 1
    put(ws, f"A{r}", "Итого: доля семей с льготой и средняя скидка", bold=True, wrap=True)
    put(ws, f"B{r}", f"=SUM(B{first}:B{last})", F_PCT, bold=True)
    put(ws, f"C{r}", None)
    put(ws, f"D{r}", f"=SUM(D{first}:D{last})", '0.00%', bold=True)
    put(ws, f"E{r}", "Средняя скидка применяется к чеку B2C во всех сценариях")
    VET["share"] = f"{q(VET_SHEET)}!$B${r}"
    VET["discount"] = f"{q(VET_SHEET)}!$D${r}"
    r += 1
    put(ws, f"A{r}", "Минут координатора в месяц на семью с льготой", wrap=True)
    put(ws, f"B{r}", M.VETERAN_PAPERWORK_MIN[0], F_NUM, color=INPUT_COLOR)
    put(ws, f"C{r}", None)
    put(ws, f"D{r}", None)
    put(ws, f"E{r}", f"{M.VETERAN_PAPERWORK_MIN[2]} [{M.VETERAN_PAPERWORK_MIN[1]}]", wrap=True)
    VET["minutes"] = f"{q(VET_SHEET)}!$B${r}"

    r += 2
    put(ws, f"A{r}", "Неценовые условия", bold=True, border=False)
    r += 1
    header(ws, r, ["Условие", "Содержание"])
    ws.merge_cells(f"B{r}:E{r}")
    for col in "CDE":
        ws[f"{col}{r}"].border = BOX
    for name, text in VET_CONDITIONS:
        r += 1
        put(ws, f"A{r}", name)
        ws.merge_cells(f"B{r}:E{r}")
        put(ws, f"B{r}", text, wrap=True)
        for col in "CDE":
            ws[f"{col}{r}"].border = BOX
        ws.row_dimensions[r].height = 30

    r += 2
    put(ws, f"A{r}", "Стоимость льгот для сервиса (недополученная выручка без НДС)", bold=True, border=False)
    r += 1
    header(ws, r, ["Сценарий", "За 60 мес., руб.", "Доля выручки за 60 мес."])
    VET["cost_rows"] = []
    for scen in SCEN_COLS:
        r += 1
        put(ws, f"A{r}", scen)
        VET["cost_rows"].append((scen, r))
    page_setup(ws)
    return ws


def fill_veteran_costs(wb):
    ws = wb[VET_SHEET]
    label = "Льготы ветеранам за 60 мес. (недополученная выручка), руб."
    for i, (scen, r) in enumerate(VET["cost_rows"], start=2):
        col = get_column_letter(i)
        put(ws, f"B{r}", f"='Сводка'!${col}${SUMMARY_ROWS[label]}", F_RUB, color=LINK_COLOR)
        rev = f"SUM({q(scen)}!${COL['rev']}${FIRST}:${COL['rev']}${LAST})"
        put(ws, f"C{r}", f"=B{r}/({rev}+B{r})", F_PCT)


# ---------------------------------------------------------------------------
# Лист «Время координатора»
# ---------------------------------------------------------------------------

TIME_TOTALS = {}


def build_time(wb):
    ws = wb.create_sheet("Время координатора")
    title(ws, "Бюджет живого времени координатора на одну семью в месяц",
          "Нагрузка на координатора рассчитывается от минут, а не задается числом")
    widths(ws, [16, 52, 16, 16, 16])
    r = 4
    header(ws, r, ["Тариф", "Операция", "Раз в месяц", "Минут на операцию", "Минут в месяц"], height=32)
    for plan in ("Базовый", "Расширенный"):
        first = r + 1
        for p, op, n, t in TIME_ROWS:
            if p != plan:
                continue
            r += 1
            put(ws, f"A{r}", p)
            put(ws, f"B{r}", op, wrap=True)
            put(ws, f"C{r}", n, F_NUM1, color=INPUT_COLOR)
            put(ws, f"D{r}", t, F_NUM, color=INPUT_COLOR)
            put(ws, f"E{r}", f"=C{r}*D{r}", F_NUM1)
        r += 1
        put(ws, f"A{r}", plan)
        put(ws, f"B{r}", VETERAN_OP, wrap=True)
        put(ws, f"C{r}", f"={VET['share']}", F_DEC, color=LINK_COLOR)
        put(ws, f"D{r}", f"={VET['minutes']}", F_NUM, color=LINK_COLOR)
        put(ws, f"E{r}", f"=C{r}*D{r}", F_NUM1)
        r += 1
        TIME_TOTALS[plan] = r
        put(ws, f"A{r}", f"Итого «{plan}»", bold=True)
        put(ws, f"B{r}", None)
        put(ws, f"C{r}", None)
        put(ws, f"D{r}", None)
        put(ws, f"E{r}", f"=SUM(E{first}:E{r - 1})", F_NUM1, bold=True)
    page_setup(ws, landscape=False)


# ---------------------------------------------------------------------------
# Лист «Бюджет MVP»
# ---------------------------------------------------------------------------

MVP_TOTAL_REF = {}


def build_mvp(wb, P):
    ws = wb.create_sheet("Бюджет MVP")
    title(ws, "Бюджет MVP на 3 месяца", "Итог по статье = количество x цена x месяцев x коэффициент взносов (для ФОТ)")
    widths(ws, [22, 50, 10, 14, 10, 10, 15, 12, 60])
    r = 4
    header(ws, r, ["Категория", "Статья", "Кол-во", "Цена за ед., руб.", "Месяцев", "Взносы", "Итого, руб.",
                   "Метка", "Комментарий"], height=32)
    pm = P.fixed("payroll_mult")
    first = r + 1
    for cat, item, qty, unit, months, payroll, tag, note in M.MVP_ITEMS:
        r += 1
        put(ws, f"A{r}", cat, wrap=True)
        put(ws, f"B{r}", item, wrap=True)
        put(ws, f"C{r}", qty, F_DEC if qty % 1 else F_NUM, color=INPUT_COLOR)
        put(ws, f"D{r}", unit, F_RUB, color=INPUT_COLOR)
        put(ws, f"E{r}", months, F_NUM1 if months % 1 else F_NUM, color=INPUT_COLOR)
        put(ws, f"F{r}", "да" if payroll else "нет", align="center", color=INPUT_COLOR)
        put(ws, f"G{r}", f'=ROUND(C{r}*D{r}*E{r}*IF(F{r}="да",{pm},1),0)', F_RUB)
        put(ws, f"H{r}", tag, align="center")
        put(ws, f"I{r}", note, wrap=True)
    last = r
    r += 1
    overhead_row = r
    put(ws, f"A{r}", "Накладные")
    put(ws, f"B{r}", "Накладные от ФОТ команды MVP")
    for col in "CDEF":
        put(ws, f"{col}{r}", None)
    put(ws, f"G{r}", f'=ROUND(SUMIF(F{first}:F{last},"да",G{first}:G{last})*{P.fixed("overhead_share")},0)', F_RUB)
    put(ws, f"H{r}", "допущение", align="center")
    put(ws, f"I{r}", "Рабочие места, HR, бухгалтерия")
    r += 1
    reserve_row = r
    put(ws, f"A{r}", "Резерв")
    put(ws, f"B{r}", "Резерв на непредвиденные расходы")
    for col in "CDEF":
        put(ws, f"{col}{r}", None)
    put(ws, f"G{r}", f"=ROUND((SUM(G{first}:G{last})+G{overhead_row})*{P.fixed('contingency')},0)", F_RUB)
    put(ws, f"H{r}", "допущение", align="center")
    put(ws, f"I{r}", M.MVP_CONTINGENCY[2])
    r += 1
    total_row = r
    put(ws, f"A{r}", "Итого бюджет MVP", bold=True)
    for col in "BCDEF":
        put(ws, f"{col}{r}", None)
    put(ws, f"G{r}", f"=SUM(G{first}:G{reserve_row})", F_RUB, bold=True)
    put(ws, f"H{r}", None)
    put(ws, f"I{r}", None)
    r += 1
    put(ws, f"A{r}", "Лимит бюджета")
    for col in "BCDEF":
        put(ws, f"{col}{r}", None)
    put(ws, f"G{r}", f"={P.fixed('mvp_limit')}", F_RUB, color=LINK_COLOR)
    put(ws, f"H{r}", None)
    put(ws, f"I{r}", "Условие кейса")
    r += 1
    put(ws, f"A{r}", "Запас до лимита", bold=True)
    for col in "BCDEF":
        put(ws, f"{col}{r}", None)
    put(ws, f"G{r}", f"=G{r - 1}-G{total_row}", F_RUB, bold=True)
    put(ws, f"H{r}", None)
    put(ws, f"I{r}", f'=IF(G{r}>=0,"Бюджет в пределах лимита","Превышение лимита")')
    MVP_TOTAL_REF["ref"] = f"{q('Бюджет MVP')}!$G${total_row}"

    r += 2
    put(ws, f"A{r}", "Бюджет по категориям", bold=True, border=False)
    r += 1
    header(ws, r, ["Категория", "Сумма, руб.", "Доля"])
    cats = []
    for item in M.MVP_ITEMS:
        if item[0] not in cats:
            cats.append(item[0])
    cats += ["Накладные", "Резерв"]
    cat_first = r + 1
    for cat in cats:
        r += 1
        put(ws, f"A{r}", cat)
        put(ws, f"B{r}", f'=SUMIF($A${first}:$A${reserve_row},A{r},$G${first}:$G${reserve_row})', F_RUB)
        put(ws, f"C{r}", f"=B{r}/$G${total_row}", F_PCT)
    r += 1
    put(ws, f"A{r}", "Итого", bold=True)
    put(ws, f"B{r}", f"=SUM(B{cat_first}:B{r - 1})", F_RUB, bold=True)
    put(ws, f"C{r}", f"=SUM(C{cat_first}:C{r - 1})", F_PCT, bold=True)
    ws.freeze_panes = "A5"
    page_setup(ws)


# ---------------------------------------------------------------------------
# Помесячные листы сценариев
# ---------------------------------------------------------------------------

FIRST = 6
LAST = FIRST + M.MONTHS - 1
YEAR_FIRST = LAST + 5
YEAR_LAST = YEAR_FIRST + M.MONTHS // 12 - 1

MONTH_COLS = [
    # ключ, заголовок, формат, ширина
    ("m", "Месяц", F_INT, 7),
    ("y", "Год", F_INT, 6),
    ("phase", "Фаза", None, 15),
    ("sal_k", "Индекс зарплат", F_DEC3, 9),
    ("price_k", "Индекс цен", F_DEC3, 9),
    ("arpu", "Средний чек B2C с НДС, руб.", F_RUB, 11),
    ("mkt", "Маркетинг, руб.", F_RUB, 12),
    ("new_int", "Новые семьи: внутренние каналы", F_NUM1, 11),
    ("new_ext", "Новые семьи: внешние каналы", F_NUM1, 11),
    ("new_vet", "Новые семьи: канал фонда «Защитники Отечества»", F_NUM1, 11),
    ("new_b2c", "Новые семьи B2C", F_NUM1, 10),
    ("b2c", "Семьи B2C", F_NUM1, 10),
    ("b2b_add", "Приток B2B2C", F_NUM1, 10),
    ("b2b", "Семьи B2B2C", F_NUM1, 10),
    ("fam", "Всего семей", F_NUM1, 10),
    ("fam_ext", "Семьи «Расши-ренный»", F_NUM1, 10),
    ("fam_basic", "Семьи «Базовый» и B2B2C", F_NUM1, 10),
    ("coords", "Координа-торы, чел.", F_NUM, 10),
    ("sups", "Супер-визоры, чел.", F_NUM, 9),
    ("rev_sub", "Выручка: подписка B2C, руб.", F_RUB, 13),
    ("rev_b2b", "Выручка: B2B2C, руб.", F_RUB, 13),
    ("rev_comm", "Выручка: комиссия партнеров, руб.", F_RUB, 12),
    ("rev", "Выручка всего, руб.", F_RUB, 13),
    ("vet_cost", "Партнерский канал фонда, руб.", F_RUB, 12),
    ("coord_pay", "ФОТ координаторов и супервизоров, руб.", F_RUB, 13),
    ("onboard", "Обучение новых координаторов, руб.", F_RUB, 12),
    ("variable", "Переменные расходы на семьи, руб.", F_RUB, 12),
    ("seats", "Рабочие места, руб.", F_RUB, 11),
    ("team", "Постоянная команда, руб.", F_RUB, 13),
    ("overhead", "Накладные, руб.", F_RUB, 12),
    ("other_fixed", "Размещение и юрлицо, руб.", F_RUB, 12),
    ("capex", "Инвестиции, руб.", F_RUB, 12),
    ("mvp", "Бюджет MVP, руб.", F_RUB, 12),
    ("cost", "Расходы всего, руб.", F_RUB, 13),
    ("op", "Операционный результат, руб.", F_RUB, 13),
    ("cum", "Накопленный результат, руб.", F_RUB, 14),
    ("cross", "Кросс-продажи (эффект группы), руб.", F_RUB, 12),
    ("benefits", "Льготы ветеранам: недополученная выручка, руб.", F_RUB, 12),
    ("tax", "Налог на прибыль, руб.", F_RUB, 12),
    ("cf", "Денежный поток после налога, руб.", F_RUB, 13),
    ("disc", "Коэффициент дисконтирования", '0.0000', 11),
    ("be_flag", "Признак безубыточности (месяц)", F_INT, 11),
    ("pb_flag", "Признак окупаемости (месяц)", F_INT, 11),
]
COL = {key: get_column_letter(i) for i, (key, *_) in enumerate(MONTH_COLS, start=1)}


def month_formulas(r, is_mvp, p, P, scen):
    """Формулы одной строки помесячного расчета. p(key) возвращает адрес параметра сценария."""
    c = {k: f"{v}{r}" for k, v in COL.items()}
    prev = {k: f"{v}{r - 1}" for k, v in COL.items()}
    m = c["m"]
    f = {
        "m": r - FIRST + 1,
        "y": f"=INT(({m}-1)/12)+1",
        "phase": "MVP" if is_mvp else "Масштабирование",
        "sal_k": f"=(1+{p('salary_index_y')})^({c['y']}-1)",
        "price_k": f"=(1+{p('price_index_y')})^({c['y']}-1)",
        "arpu": f"={p('arpu_gross')}*{c['price_k']}",
        "b2b_add": 0 if is_mvp else
        f"=IF({m}>={p('b2b_start_month')},{p('cohort_b2b')}/{p('b2b_ramp_months')}"
        f"*MIN({p('b2b_ramp_months')},{m}-{p('b2b_start_month')}+1),0)",
        "fam": f"={c['b2c']}+{c['b2b']}",
        "fam_ext": f"={c['b2c']}*{p('share_ext')}",
        "fam_basic": f"={c['b2c']}*(1-{p('share_ext')})+{c['b2b']}",
        "rev_b2b": f"={c['b2b']}*{p('b2b_price')}*{c['price_k']}/(1+{p('vat')})",
        "rev_comm": f"={c['fam']}*{p('partner_spend')}*{p('partner_commission')}/(1+{p('vat')})",
        "rev": f"={c['rev_sub']}+{c['rev_b2b']}+{c['rev_comm']}",
        "benefits": f"={c['rev_sub']}*{p('vet_k')}",
        "cost": f"={c['mkt']}+SUM({c['vet_cost']}:{c['mvp']})",
        "op": f"={c['rev']}-{c['cost']}",
        "cum": f"={c['op']}" if r == FIRST else f"={prev['cum']}+{c['op']}",
        "tax": f"=IF(MOD({m},12)=0,INDEX($F${YEAR_FIRST}:$F${YEAR_LAST},{c['y']}),0)",
        "cf": f"={c['op']}-{c['tax']}",
        "disc": f"=1/(1+{p('disc_m')})^{m}",
        "be_flag": f'=IF(AND({m}>{p("mvp_months")},MIN({COL["op"]}{r}:{COL["op"]}${LAST})>0),{m},"")',
        "pb_flag": f'=IF(MIN({COL["cum"]}{r}:{COL["cum"]}${LAST})>0,{m},"")',
    }
    if is_mvp:
        month = r - FIRST + 1
        fam_now = f"{A_Q}!$B${P.mvp_fam[month]}"
        fam_prev = f"{A_Q}!$B${P.mvp_fam[month - 1]}" if month > 1 else "0"
        zero = ["mkt", "new_int", "new_ext", "new_vet", "vet_cost", "b2b", "sups", "coord_pay", "onboard", "variable", "seats", "team",
                "overhead", "other_fixed", "capex", "cross"]
        f.update({k: 0 for k in zero})
        f.update({
            "new_b2c": f"={fam_now}-{fam_prev}",
            "b2c": f"={fam_now}",
            "coords": f"=IF({m}>1,{p('min_coords')},0)",
            "rev_sub": f"={c['b2c']}*{p('mvp_paying_share')}*{c['arpu']}/(1+{p('vat')})",
            "mvp": f"={MVP_TOTAL_REF['ref']}/{p('mvp_months')}",
        })
        return f
    first_scale = r == FIRST + M.MVP_MONTHS
    team = (f"SUMPRODUCT({A_Q}!$B${P.team_first}:$B${P.team_last},{A_Q}!$C${P.team_first}:$C${P.team_last},"
            f"--({A_Q}!$D${P.team_first}:$D${P.team_last}<={m}))")
    f.update({
        "mkt": f"={p('mkt_start')}" if first_scale else f"=MIN({prev['mkt']}*(1+{p('mkt_growth_m')}),{p('mkt_cap')})",
        "new_int": f"=MIN({p('int_cap')},{c['mkt']}/{p('int_cac')})",
        "new_ext": f"=MAX(0,{c['mkt']}-{c['new_int']}*{p('int_cac')})/{p('ext_cac')}",
        "new_vet": f"=IF({m}>={p('vet_start_month')},{p('vet_new')},0)",
        "vet_cost": f"={c['new_vet']}*{p('vet_cac')}",
        "new_b2c": f"={c['new_int']}+{c['new_ext']}+{c['new_vet']}",
        "b2c": f"=({prev['b2c']}-{prev['new_b2c']})*(1-{p('churn_b2c')})+{prev['new_b2c']}*(1-{p('churn_first')})"
               f"+{c['new_b2c']}",
        "b2b": f"={prev['b2b']}*(1-{p('churn_b2b')})+{c['b2b_add']}",
        "coords": f"=MAX({p('min_coords')},ROUNDUP({c['fam_basic']}/{p('load_b')}+{c['fam_ext']}/{p('load_e')},0))",
        "sups": f"=ROUNDUP({c['coords']}/{p('supervisor_ratio')},0)",
        "rev_sub": f"={c['b2c']}*{c['arpu']}/(1+{p('vat')})",
        "coord_pay": f"=({c['coords']}*{p('coord_salary')}+{c['sups']}*{p('supervisor_salary')})*{c['sal_k']}"
                     f"*{p('payroll_mult')}",
        "onboard": f"=(MAX(0,{c['coords']}-{prev['coords']})+{c['coords']}*{p('coord_turnover_y')}/12)"
                   f"*{p('onboarding_cost')}*{c['sal_k']}",
        "variable": f"={c['fam']}*({p('telemed_per_family')}+{p('telecom_per_family')})"
                    f"+({c['b2c']}*{c['arpu']}+{c['b2b']}*{p('b2b_price')}*{c['price_k']})*{p('acquiring')}",
        "seats": f"=({c['coords']}+{c['sups']})*{p('seat_cost')}",
        "team": f"={team}*{c['sal_k']}*{p('payroll_mult')}",
        "overhead": f"=({c['team']}+{c['coord_pay']})*{p('overhead_share')}",
        "other_fixed": f"={p('hosting_base')}+{p('hosting_per_1000')}*{c['fam']}/1000+{p('legal_entity_m')}",
        "capex": f"=SUMIF({A_Q}!$B${P.capex_first}:$B${P.capex_last},{m},"
                 f"{A_Q}!$C${P.capex_first}:$C${P.capex_last})",
        "mvp": 0,
        "cross": f"={c['fam']}*{p('cross_rate_m')}*{p('cross_margin')}",
    })
    return f


def build_month_sheet(wb, P, scen):
    ws = wb.create_sheet(scen)
    label = "сценарий «Версия 1 без триажа»" if scen == "Версия 1" else f"сценарий «{scen}»"
    title(ws, f"Помесячный расчет: {label}",
          "Суммы в рублях без НДС. Месяцы 1-3: MVP. Налог на прибыль начисляется в последнем месяце года.")
    p = lambda key: P.ref(key, scen)
    header(ws, FIRST - 1, [h for _, h, _, _ in MONTH_COLS], height=62)
    for i, (*_, w) in enumerate(MONTH_COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(FIRST, LAST + 1):
        is_mvp = r - FIRST + 1 <= M.MVP_MONTHS
        f = month_formulas(r, is_mvp, p, P, scen)
        for key, _, fmt, _ in MONTH_COLS:
            put(ws, f"{COL[key]}{r}", f[key], fmt, align="center" if key in ("m", "y", "phase") else None)

    # Годовые итоги и налог с переносом убытков
    r = YEAR_FIRST - 2
    put(ws, f"A{r}", "Годовые итоги и налог на прибыль с переносом убытков", bold=True, border=False)
    header(ws, YEAR_FIRST - 1, ["Год", "Выручка, руб.", "Операционный результат, руб.",
                                "Убыток к переносу на начало, руб.", "Зачтено убытка, руб.",
                                "Налог на прибыль, руб.", "Убыток к переносу на конец, руб.",
                                "Кросс-продажи, руб."], height=48)
    years = f"$B${FIRST}:$B${LAST}"
    for i, r in enumerate(range(YEAR_FIRST, YEAR_LAST + 1), start=1):
        put(ws, f"A{r}", i, F_INT, align="center")
        put(ws, f"B{r}", f"=SUMIF({years},A{r},${COL['rev']}${FIRST}:${COL['rev']}${LAST})", F_RUB)
        put(ws, f"C{r}", f"=SUMIF({years},A{r},${COL['op']}${FIRST}:${COL['op']}${LAST})", F_RUB)
        put(ws, f"D{r}", 0 if i == 1 else f"=G{r - 1}", F_RUB)
        put(ws, f"E{r}", f"=IF(C{r}>0,MIN(D{r},C{r}*{p('loss_carry_limit')}),0)", F_RUB)
        put(ws, f"F{r}", f"=IF(C{r}>0,(C{r}-E{r})*{p('profit_tax')},0)", F_RUB)
        put(ws, f"G{r}", f"=D{r}+IF(C{r}<0,-C{r},0)-E{r}", F_RUB)
        put(ws, f"H{r}", f"=SUMIF({years},A{r},${COL['cross']}${FIRST}:${COL['cross']}${LAST})", F_RUB)
    r = YEAR_LAST + 1
    put(ws, f"A{r}", "Итого", bold=True)
    for col in "BCH":
        put(ws, f"{col}{r}", f"=SUM({col}{YEAR_FIRST}:{col}{YEAR_LAST})", F_RUB, bold=True)
    put(ws, f"F{r}", f"=SUM(F{YEAR_FIRST}:F{YEAR_LAST})", F_RUB, bold=True)
    for col in "DEG":
        put(ws, f"{col}{r}", None)
    ws.freeze_panes = f"D{FIRST}"
    page_setup(ws)
    ws.print_title_rows = f"{FIRST - 1}:{FIRST - 1}"


# ---------------------------------------------------------------------------
# Сводка по сценариям
# ---------------------------------------------------------------------------

SUMMARY_ROWS = {}


def build_summary(wb):
    ws = wb.create_sheet("Сводка")
    title(ws, "Сводные результаты по сценариям", "Горизонт 60 месяцев, NPV по месячной ставке дисконтирования")
    widths(ws, [52, 18, 18, 18, 18])
    r = 4
    header(ws, r, ["Показатель", "Пессимистичный", "Базовый", "Оптимистичный", "Версия 1 без триажа"], height=32)
    s36 = FIRST + M.SHORT - 1

    def rng(key, last=LAST):
        return f"${COL[key]}${FIRST}:${COL[key]}${last}"

    rows = [
        ("NPV до налога за 36 мес., руб.", lambda s: f"=SUMPRODUCT({s}!{rng('op', s36)},{s}!{rng('disc', s36)})", F_RUB),
        ("NPV после налога за 36 мес., руб.", lambda s: f"=SUMPRODUCT({s}!{rng('cf', s36)},{s}!{rng('disc', s36)})", F_RUB),
        ("NPV до налога за 60 мес., руб.", lambda s: f"=SUMPRODUCT({s}!{rng('op')},{s}!{rng('disc')})", F_RUB),
        ("NPV после налога за 60 мес., руб.", lambda s: f"=SUMPRODUCT({s}!{rng('cf')},{s}!{rng('disc')})", F_RUB),
        ("NPV до налога за 60 мес. с кросс-продажами группы, руб.",
         lambda s: f"=SUMPRODUCT({s}!{rng('op')}+{s}!{rng('cross')},{s}!{rng('disc')})", F_RUB),
        ("Пиковая потребность в деньгах, руб.", lambda s: f"=MAX(0,-MIN({s}!{rng('cum')}))", F_RUB),
        ("Месяц операционной безубыточности",
         lambda s: f'=IF(COUNT({s}!{rng("be_flag")})=0,"не достигается",MIN({s}!{rng("be_flag")}))', F_INT),
        ("Месяц окупаемости (накопленный результат > 0)",
         lambda s: f'=IF(COUNT({s}!{rng("pb_flag")})=0,"не достигается",MIN({s}!{rng("pb_flag")}))', F_INT),
        ("Семей на конец 36 мес.", lambda s: f"={s}!${COL['fam']}${s36}", F_NUM),
        ("Семей на конец 60 мес.", lambda s: f"={s}!${COL['fam']}${LAST}", F_NUM),
        ("Координаторов на конец 60 мес.", lambda s: f"={s}!${COL['coords']}${LAST}", F_NUM),
        ("Выручка за MVP, руб.", lambda s: f"=SUM({s}!${COL['rev']}${FIRST}:${COL['rev']}${FIRST + M.MVP_MONTHS - 1})",
         F_RUB),
    ]
    for y in range(M.MONTHS // 12):
        rows.append((f"Выручка, год {y + 1}, руб.", lambda s, y=y: f"={s}!$B${YEAR_FIRST + y}", F_RUB))
    for y in range(M.MONTHS // 12):
        rows.append((f"Операционный результат, год {y + 1}, руб.", lambda s, y=y: f"={s}!$C${YEAR_FIRST + y}", F_RUB))
    rows.append(("Кросс-продажи за 60 мес. (эффект группы), руб.", lambda s: f"={s}!$H${YEAR_LAST + 1}", F_RUB))
    rows.append(("Льготы ветеранам за 60 мес. (недополученная выручка), руб.",
                 lambda s: f"=SUM({s}!{rng('benefits')})", F_RUB))

    for label, fn, fmt in rows:
        r += 1
        SUMMARY_ROWS[label] = r
        put(ws, f"A{r}", label, wrap=True)
        for i, scen in enumerate(SCEN_COLS, start=2):
            put(ws, f"{get_column_letter(i)}{r}", fn(q(scen)), fmt, align="right")
    ws.freeze_panes = "B5"
    page_setup(ws, landscape=False)


# ---------------------------------------------------------------------------
# Юнит-экономика
# ---------------------------------------------------------------------------

def build_unit(wb, P):
    ws = wb.create_sheet("Юнит-экономика")
    title(ws, "Юнит-экономика одной семьи B2C", "Устоявшаяся семья, первый год, без индексации цен и зарплат")
    widths(ws, [56, 18, 18, 18, 18])
    r = 4
    header(ws, r, ["Показатель", "Пессимистичный", "Базовый", "Оптимистичный", "Версия 1 без триажа"], height=32)
    names = {}
    rows = [
        ("arpu_l", "Средний чек по прейскуранту с НДС, руб./мес.", lambda a: f"={a('arpu_list')}", F_RUB),
        ("vet_d", "Средняя скидка по льготам ветеранам", lambda a: f"={VET['discount']}", '0.0%'),
        ("arpu_g", "Средний чек с НДС с учетом льгот, руб./мес.", lambda a: "={arpu_l}*(1-{vet_d})", F_RUB),
        ("arpu_n", "Средний чек без НДС, руб./мес.", lambda a: f"={{arpu_g}}/(1+{a('vat')})", F_RUB),
        ("comm", "Комиссия партнеров без НДС, руб./мес.",
         lambda a: f"={a('partner_spend')}*{a('partner_commission')}/(1+{a('vat')})", F_RUB),
        ("lb", "Семей «Базовый» на координатора", lambda a: f"={a('load_b')}", F_NUM),
        ("le", "Семей «Расширенный» на координатора", lambda a: f"={a('load_e')}", F_NUM),
        ("units", "Доля ставки координатора на семью",
         lambda a: f"=(1-{a('share_ext')})/{{lb}}+{a('share_ext')}/{{le}}", '0.0000'),
        ("per_coord", "Полная стоимость координатора с супервизией, руб./мес.",
         lambda a: f"={a('coord_salary')}*{a('payroll_mult')}*(1+{a('overhead_share')})"
                   f"+{a('supervisor_salary')}*{a('payroll_mult')}*(1+{a('overhead_share')})/{a('supervisor_ratio')}"
                   f"+{a('seat_cost')}*(1+1/{a('supervisor_ratio')})"
                   f"+{a('onboarding_cost')}*{a('coord_turnover_y')}/12", F_RUB),
        ("coord", "Координатор, супервизия, накладные на семью, руб./мес.", lambda a: "={per_coord}*{units}", F_RUB),
        ("var", "Прочие переменные на семью, руб./мес.",
         lambda a: f"={a('telemed_per_family')}+{a('telecom_per_family')}+{{arpu_g}}*{a('acquiring')}"
                   f"+{a('hosting_per_1000')}/1000", F_RUB),
        ("contrib", "Маржинальный доход на семью, руб./мес.", lambda a: "={arpu_n}+{comm}-{coord}-{var}", F_RUB),
        ("margin", "Маржинальность", lambda a: "={contrib}/({arpu_n}+{comm})", F_PCT),
        ("life", "Средний срок жизни, мес.", lambda a: f"=1+(1-{a('churn_first')})/{a('churn_b2c')}", F_NUM1),
        ("ltv", "LTV, руб.", lambda a: "={contrib}*{life}", F_RUB),
        ("cac", "CAC при полном бюджете (смесь каналов), руб.",
         lambda a: f"={a('mkt_cap')}/({a('int_cap')}+({a('mkt_cap')}-{a('int_cap')}*{a('int_cac')})/{a('ext_cac')})",
         F_RUB),
        ("ltv_cac", "LTV/CAC", lambda a: "={ltv}/{cac}", '0.00"x";(0.00"x")'),
        ("payback", "Окупаемость CAC, мес.", lambda a: '=IF({contrib}>0,{cac}/{contrib},"не окупается")', F_NUM1),
    ]
    for key, *_ in rows:
        r += 1
        names[key] = r
    r = 4
    for key, label, fn, fmt in rows:
        r += 1
        put(ws, f"A{r}", label, wrap=True)
        for i, scen in enumerate(SCEN_COLS, start=2):
            col = get_column_letter(i)
            formula = fn(lambda k: P.ref(k, scen)).format(**{k: f"{col}{v}" for k, v in names.items()})
            put(ws, f"{col}{r}", formula, fmt, align="right")
    page_setup(ws, landscape=False)
    return names


# ---------------------------------------------------------------------------
# Рынок
# ---------------------------------------------------------------------------

def build_market(wb, P):
    ws = wb.create_sheet("Рынок")
    title(ws, "Оценка размера рынка (порядок величины)", "Базовый сценарий")
    widths(ws, [64, 20, 14, 64])
    r = 4
    header(ws, r, ["Показатель", "Значение", "Метка", "Источник или обоснование"])
    refs = []
    for name, val, tag, src in M.MARKET:
        r += 1
        refs.append(f"B{r}")
        put(ws, f"A{r}", name, wrap=True)
        put(ws, f"B{r}", val, F_PCT if val < 1 else F_NUM, color=INPUT_COLOR)
        put(ws, f"C{r}", tag, align="center")
        put(ws, f"D{r}", src, wrap=True)
    n65, chronic, separate, coord, sam_share = refs
    arpu = P.ref("arpu_gross", "Базовый")
    som = f"'Сводка'!$C${SUMMARY_ROWS['Семей на конец 60 мес.']}"
    calc = [
        ("TAM, семей", f"={n65}*{chronic}*{separate}*{coord}", F_NUM, "Произведение долей"),
        ("TAM, руб. в год", "={tam}*{arpu}*12", F_RUB, "TAM x средний чек с НДС x 12"),
        ("SAM, семей", f"={{tam}}*{sam_share}", F_NUM, "TAM x доля платежеспособных в городах 250 тыс.+"),
        ("SAM, руб. в год", "={sam}*{arpu}*12", F_RUB, "SAM x средний чек с НДС x 12"),
        ("SOM на 60 мес., семей (базовый сценарий)", f"={som}", F_NUM, "Лист «Сводка»"),
        ("Доля SOM от SAM", "={som_r}/{sam}", '0.00%', "SOM / SAM"),
    ]
    cells = {}
    for label, formula, fmt, note in calc:
        r += 1
        key = {"TAM, семей": "tam", "SAM, семей": "sam"}.get(label)
        if key:
            cells[key] = f"B{r}"
        if label.startswith("SOM"):
            cells["som_r"] = f"B{r}"
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", formula.format(arpu=arpu, **cells), fmt, color=LINK_COLOR if label.startswith("SOM") else "000000")
        put(ws, f"C{r}", "расчет", align="center")
        put(ws, f"D{r}", note)
    page_setup(ws)


# ---------------------------------------------------------------------------
# Чувствительность, Монте-Карло, риски
# ---------------------------------------------------------------------------

STOCHASTIC_NOTE = ("Результаты получены многократным прогоном модели в model.py (каждая строка требует отдельного "
                   "пересчета всех 60 месяцев), поэтому перенесены значениями. Пересчет: python model.py.")


def build_sensitivity(wb, base_npv, tornado):
    ws = wb.create_sheet("Чувствительность")
    title(ws, "Чувствительность NPV до налога за 60 мес. к параметрам (базовый сценарий)", STOCHASTIC_NOTE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 30
    widths(ws, [40, 20, 20, 20, 20])
    put(ws, "A4", "NPV базового сценария, руб.", bold=True)
    put(ws, "B4", base_npv, F_RUB, bold=True)
    header(ws, 6, ["Параметр", "NPV при -25%, руб.", "NPV при +25%, руб.", "Размах, руб.", "Ранг"])
    r = 6
    for i, row in enumerate(tornado.itertuples(index=False), start=1):
        r += 1
        put(ws, f"A{r}", row[0])
        put(ws, f"B{r}", row[1], F_RUB)
        put(ws, f"C{r}", row[2], F_RUB)
        put(ws, f"D{r}", f"=ABS(C{r}-B{r})", F_RUB)
        put(ws, f"E{r}", i, F_INT, align="center")
    r += 2
    put(ws, f"A{r}", "Для параметра «Цены тарифов» одновременно меняются цены обоих тарифов и опции B2B2C.",
        border=False, italic=True)
    page_setup(ws, landscape=False)


MC_LABELS = {
    "ext_cac": "CAC во внешних каналах, руб.", "int_cap": "Емкость внутренних каналов, семей/мес.",
    "churn_first": "Отток после первого месяца", "churn_b2c": "Отток B2C в месяц",
    "time_mult": "Множитель минут координатора", "share_ext": "Доля тарифа «Расширенный»",
    "coord_salary": "Зарплата координатора, руб.", "b2b_activation": "Активация B2B2C",
    "b2b_clients_per_month": "Новых корпоративных клиентов в месяц", "partner_spend": "Покупки услуг партнеров, руб./мес.",
}


def build_monte_carlo(wb, mc):
    ws = wb.create_sheet("Монте-Карло")
    title(ws, "Анализ рисков методом Монте-Карло", STOCHASTIC_NOTE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 30
    widths(ws, [46, 18, 18, 18])
    put(ws, "A4", "Треугольные распределения параметров", bold=True, border=False)
    header(ws, 5, ["Параметр", "Минимум", "Наиболее вероятное", "Максимум"])
    r = 5
    for key, (lo, mode, hi) in M.MC_DIST.items():
        r += 1
        fmt = F_PCT if hi < 1 else (F_DEC if hi < 10 else F_NUM)
        put(ws, f"A{r}", MC_LABELS[key])
        for col, v in zip("BCD", (lo, mode, hi)):
            put(ws, f"{col}{r}", v, fmt, color=INPUT_COLOR)
    r += 2
    put(ws, f"A{r}", "Результаты", bold=True, border=False)
    r += 1
    header(ws, r, ["Показатель", "Значение"])
    for k, v in mc.items():
        r += 1
        is_pct = "%" in k
        put(ws, f"A{r}", k.replace(", %", "").replace(", руб", ", руб."), wrap=True)
        put(ws, f"B{r}", v / 100 if is_pct else v, F_PCT if is_pct else F_NUM, align="right")
    page_setup(ws, landscape=False)


def build_risks(wb):
    ws = wb.create_sheet("Риски модели")
    title(ws, "Подводные камни и их учет в модели")
    widths(ws, [44, 54, 60, 12])
    header(ws, 3, ["Риск", "Суть", "Как учтено", "Влияние"])
    for r, row in enumerate(M.PITFALLS, start=4):
        for col, v in zip("ABCD", row):
            put(ws, f"{col}{r}", v, wrap=True, align="center" if col == "D" else None)
    page_setup(ws)


def build_cover(wb):
    ws = wb.active
    ws.title = "Титул"
    widths(ws, [4, 34, 90])
    ws["B2"].value = "Финансовая модель продукта «ИНГОССЕМЬЯ»"
    ws["B2"].font = font(bold=True, size=16)
    ws["B3"].value = "Кейс IngoSight 2026. Все суммы в рублях; выручка и затраты без НДС, если не указано иное."
    ws["B3"].font = font(italic=True)
    put(ws, "B5", "Лист", bold=True)
    put(ws, "C5", "Содержание", bold=True)
    sheets = [
        ("Допущения", "Все исходные параметры по четырем сценариям, расчетные параметры, команда, инвестиции"),
        ("Льготы ветеранам", "Особые условия для ветеранов СВО и ветеранов: скидки, неценовые условия, стоимость программы"),
        ("Время координатора", "Бюджет минут координатора на семью и итог по тарифам"),
        ("Бюджет MVP", "Смета MVP на 3 месяца, проверка лимита 5 млн руб., итоги по категориям"),
        ("Сводка", "NPV, окупаемость, пиковая потребность, выручка и результат по годам для всех сценариев"),
        ("Юнит-экономика", "Маржинальный доход, LTV, CAC и окупаемость CAC на одну семью"),
        ("Рынок", "TAM, SAM и SOM"),
        ("Пессимистичный", "Помесячный расчет на 60 месяцев и налог по годам"),
        ("Базовый", "Помесячный расчет на 60 месяцев и налог по годам"),
        ("Оптимистичный", "Помесячный расчет на 60 месяцев и налог по годам"),
        ("Версия 1", "Первая версия продукта без голосового робота: показывает, почему продукт пересобран"),
        ("Чувствительность", "Влияние изменения параметров на ±25% на NPV"),
        ("Монте-Карло", "5000 прогонов модели со случайными параметрами"),
        ("Риски модели", "Подводные камни и как они учтены"),
    ]
    r = 5
    for name, desc in sheets:
        r += 1
        c = put(ws, f"B{r}", name)
        c.hyperlink = f"#'{name}'!A1"
        put(ws, f"C{r}", desc, wrap=True)
    r += 2
    put(ws, f"B{r}", "Обозначения", bold=True, border=False)
    for text, color in [("Синий шрифт: исходные данные и допущения, их можно менять", INPUT_COLOR),
                        ("Черный шрифт: формулы", "000000"),
                        ("Зеленый шрифт: ссылка на другой лист", LINK_COLOR),
                        ("Отрицательные значения показаны в скобках, нулевые прочерком", "000000"),
                        ("Метки у чисел: факт (из источника), расчет (на основе фактов), допущение (команды)", "000000")]:
        r += 1
        put(ws, f"B{r}", text, color=color, border=False)
        ws.merge_cells(f"B{r}:C{r}")
    page_setup(ws, landscape=False)


def build(path, tornado_base, tornado, mc):
    wb = Workbook()
    P = Params()
    build_cover(wb)
    build_veterans(wb)  # нужен до бюджета минут и допущений: адреса доли и скидки
    build_time(wb)  # нужен до допущений: адреса итогов минут
    build_assumptions(wb, P)
    build_mvp(wb, P)
    for scen in SCEN_COLS:
        build_month_sheet(wb, P, scen)
    build_summary(wb)
    fill_veteran_costs(wb)
    build_unit(wb, P)
    build_market(wb, P)
    build_sensitivity(wb, tornado_base, tornado)
    build_monte_carlo(wb, mc)
    build_risks(wb)
    order = ["Титул", "Допущения", VET_SHEET, "Время координатора", "Бюджет MVP", "Сводка", "Юнит-экономика", "Рынок",
             *SCEN_COLS, "Чувствительность", "Монте-Карло", "Риски модели"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.save(path)
